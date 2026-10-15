# Файлы сценария для tests/script-e2e.js: .docx, PDF с текстом, скан и старый .doc.
# Запуск: python3 tests/script-fixtures.py <папка>
import sys, os
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
out = sys.argv[1]; os.makedirs(out, exist_ok=True)
SC = [('1. ИНТ. КВАРТИРА 5А — ДЕНЬ', ['Анна стоит у окна.', 'АННА', 'Он не придёт.']),
      ('2. НАТ. ДВОР ШКОЛЫ — ВЕЧЕР', ['Дети расходятся. Фонари & тени.']),
      ('12А. ИНТ. КОРИДОР — НОЧЬ', ['Свет из-под двери.'])]
d = Document(); d.add_paragraph('ТЕСТОВЫЙ СЦЕНАРИЙ')
for h, body in SC:
    d.add_paragraph(h)
    for b in body: d.add_paragraph(b)
p = d.add_paragraph('Сноска'); p.add_run().add_break(); p.add_run('после переноса')
d.save(os.path.join(out, 'scenario.docx'))
pdfmetrics.registerFont(TTFont('DV', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
c = canvas.Canvas(os.path.join(out, 'scenario.pdf'))
for i, (h, body) in enumerate(SC):
    c.setFont('DV', 12); y = 780
    # заголовок двумя кусками, как его отдают программы вёрстки сценариев
    num, rest = h.split(' ', 1)
    c.drawString(72, y, num); c.drawString(72 + c.stringWidth(num + ' ', 'DV', 12), y, rest)
    for b in body:
        y -= 20; c.drawString(110, y, b)
    c.drawString(500, 40, str(i + 1) + '.')   # номер страницы
    c.showPage()
c.save()
c = canvas.Canvas(os.path.join(out, 'scan.pdf')); c.rect(50, 50, 400, 600, fill=1); c.showPage(); c.save()
open(os.path.join(out, 'old.doc'), 'wb').write(bytes([0xd0, 0xcf, 0x11, 0xe0, 0xa1, 0xb1, 0x1a, 0xe1]) + b'\0' * 504)
